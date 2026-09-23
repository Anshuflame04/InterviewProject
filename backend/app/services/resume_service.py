import io
import logging
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader

from app.core.config import settings
from app.db.firestore import get_user_resumes_collection
from app.prompts.resume import build_resume_extraction_prompt
from app.schemas.resume import ResumeData, ResumeResponse
from app.services.llm_service import llm_service


logger = logging.getLogger(__name__)


class ResumeService:
    """
    Handles resume extraction, parsing, and persistence.

    Raw text extraction is performed locally.
    The LLM is used only for semantic structuring
    of the extracted resume content.
    """

    ALLOWED_EXTENSIONS = {".pdf", ".docx"}

    async def extract_text(
        self,
        file: UploadFile,
    ) -> str:
        """
        Extract raw text from a PDF or DOCX file.
        """

        filename = file.filename or ""
        extension = Path(filename).suffix.lower()

        if extension not in self.ALLOWED_EXTENSIONS:
            raise ValueError(
                "Unsupported resume format. "
                "Only PDF and DOCX files are allowed."
            )

        content = await file.read()

        max_size = settings.max_resume_size_mb * 1024 * 1024

        if len(content) > max_size:
            raise ValueError(
                f"Resume exceeds the maximum size of "
                f"{settings.max_resume_size_mb} MB."
            )

        if not content:
            raise ValueError("Uploaded resume is empty.")

        if extension == ".pdf":
            text = self._extract_pdf_text(content)
        else:
            text = self._extract_docx_text(content)

        text = self._clean_text(text)

        if not text:
            raise ValueError(
                "Could not extract readable text from the resume."
            )

        return text

    @staticmethod
    def _extract_pdf_text(content: bytes) -> str:
        """
        Extract text from PDF pages.
        """

        try:
            reader = PdfReader(io.BytesIO(content))

            pages = []

            for page in reader.pages:
                page_text = page.extract_text() or ""

                if page_text.strip():
                    pages.append(page_text)

            return "\n".join(pages)

        except Exception as exc:
            logger.exception(
                "Failed to extract PDF text."
            )

            raise ValueError(
                "Unable to read the PDF file."
            ) from exc

    @staticmethod
    def _extract_docx_text(content: bytes) -> str:
        """
        Extract paragraphs and table content from DOCX.
        """

        try:
            document = Document(io.BytesIO(content))

            parts: list[str] = []

            for paragraph in document.paragraphs:
                text = paragraph.text.strip()

                if text:
                    parts.append(text)

            for table in document.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]

                    row_text = " | ".join(
                        cell
                        for cell in cells
                        if cell
                    )

                    if row_text:
                        parts.append(row_text)

            return "\n".join(parts)

        except Exception as exc:
            logger.exception(
                "Failed to extract DOCX text."
            )

            raise ValueError(
                "Unable to read the DOCX file."
            ) from exc

    @staticmethod
    def _clean_text(text: str) -> str:
        """
        Normalize extracted text without changing its meaning.
        """

        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")

        text = re.sub(
            r"[ \t]+",
            " ",
            text,
        )

        text = re.sub(
            r"\n{3,}",
            "\n\n",
            text,
        )

        return text.strip()

    async def parse_resume_text(
        self,
        resume_text: str,
    ) -> ResumeData:
        """
        Convert raw resume text into structured resume data
        using the configured LLM.
        """

        if not resume_text.strip():
            raise ValueError(
                "Resume text cannot be empty."
            )

        prompt = build_resume_extraction_prompt(
            resume_text=resume_text
        )

        return await llm_service.generate_structured(
            prompt=prompt,
            response_model=ResumeData,
            temperature=0.1,
        )

    async def process_and_save_resume(
        self,
        user_id: str,
        file: UploadFile,
    ) -> ResumeResponse:
        """
        Complete resume processing pipeline:

        Upload
            ↓
        Local text extraction
            ↓
        LLM semantic parsing
            ↓
        Firestore
        """

        resume_text = await self.extract_text(file)

        resume_data = await self.parse_resume_text(
            resume_text
        )

        collection = get_user_resumes_collection(
            user_id
        )

        document = collection.document()

        now = datetime.now(timezone.utc)

        resume_document = {
            "resume_id": document.id,
            "user_id": user_id,
            "filename": file.filename or "resume",
            "resume_text": resume_text,
            "data": resume_data.model_dump(
                mode="json"
            ),
            "created_at": now,
            "updated_at": now,
        }

        document.set(resume_document)

        return ResumeResponse(
            resume_id=document.id,
            user_id=user_id,
            filename=file.filename or "resume",
            parsed_data=resume_data,
        )

    def get_resume(
        self,
        user_id: str,
        resume_id: str,
    ) -> ResumeResponse | None:
        """
        Retrieve a resume belonging to the authenticated user.
        """

        document = (
            get_user_resumes_collection(user_id)
            .document(resume_id)
            .get()
        )

        if not document.exists:
            return None

        data = document.to_dict()

        if not data:
            return None

        stored_resume_data = data.get(
            "data",
            {},
        )

        return ResumeResponse(
            resume_id=document.id,
            user_id=user_id,
            filename=data.get(
                "filename",
                "resume",
            ),
            parsed_data=ResumeData.model_validate(
                stored_resume_data
            ),
        )

    def get_resume_collection(
        self,
        user_id: str,
    ):
        """
        Return the authenticated user's resume collection.
        """

        return get_user_resumes_collection(
            user_id
        )


resume_service = ResumeService()